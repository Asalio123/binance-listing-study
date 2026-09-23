"""Сборка BakhtievSaid_paper.docx из paper/cjsj_submission.md поверх официального
шаблона CJSJ (CJSJ-Original-Research-Template.docx).

Шаблон сохраняет sectPr (2 колонки, поля 0.75", Letter) и стили (Title 16pt,
Heading1 small-caps center, Heading2 italic, Normal TNR 10). Спец-параграфы
(abstract 9pt bold-italic с префиксом «Abstract—», Key Terms, captions 8pt,
references 8pt, author 11pt center) форматируются напрямую по разбору XML
шаблона (Угол 17 досье).

Флаги:
  --nofigs  собрать вариант без фигур и references (замер лимита 2-3 стр.:
            лимит их исключает, таблицы и текст считаются).

Запуск: /tmp/cjsj-build/bin/python publish/build_submission_docx.py [--nofigs]
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
MD = ROOT / "paper" / "cjsj_submission.md"
TPL = Path.home() / "РЕСЕРЧ" / "repro_cjsj_v1" / "template" / "CJSJ-Original-Research-Template.docx"
OUT_DIR = ROOT / "publish" / "cjsj_package"

NOFIGS = "--nofigs" in sys.argv
OUT = OUT_DIR / ("BakhtievSaid_paper_bodyonly.docx" if NOFIGS else "BakhtievSaid_paper.docx")

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*)")


def add_runs(par, text, size=None, bold=False, italic=False):
    """Добавляет runs с мини-парсером **bold** / *italic*."""
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        b, i, t = bold, italic, chunk
        if chunk.startswith("**") and chunk.endswith("**"):
            b, t = True, chunk[2:-2]
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            i, t = True, chunk[1:-1]
        r = par.add_run(t)
        r.bold, r.italic = b, i
        if size:
            r.font.size = size
    return par


def para(doc, text, style=None, align=None, size=None, bold=False, italic=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    add_runs(p, text, size=size, bold=bold, italic=italic)
    return p


def set_cell(cell, text, size, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.paragraphs[0].alignment = align
    add_runs(cell.paragraphs[0], text, size=size, bold=bold)


def table_borders(table):
    """Booktabs: верх/низ 1pt, подчёркивание заголовка 0.6pt, без вертикалей."""
    tbl = table._tbl
    rows = tbl.findall(qn("w:tr"))
    for i, row in enumerate(rows):
        for tc in row.findall(qn("w:tc")):
            tcPr = tc.find(qn("w:tcPr"))
            if tcPr is None:
                tcPr = OxmlElement("w:tcPr")
                tc.insert(0, tcPr)
            borders = OxmlElement("w:tcBorders")
            for edge, sz in (("top", 8 if i == 0 else 0),
                             ("bottom", 8 if i == len(rows) - 1
                              else (5 if i == 0 else 0))):
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:val"), "single" if sz else "none")
                el.set(qn("w:sz"), str(sz))
                borders.append(el)
            tcPr.append(borders)


def fix_widths(table, widths_in):
    """Фиксированный layout + явные tblW/tblGrid/tcW: иначе LibreOffice берёт
    равные gridCol на всю страницу и клипает таблицу по краю колонки."""
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(int(sum(widths_in) * 1440)))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    grid = tbl.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_in):
        gc.set(qn("w:w"), str(int(w * 1440)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_in):
            cell.width = Inches(w)


def main():
    doc = Document(str(TPL))
    # очистка тела шаблона, sectPr сохраняем
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    lines = MD.read_text(encoding="utf-8").splitlines()
    i = 0
    first_h1 = True
    in_refs = False
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line.strip():
            continue
        if line.startswith("AUTHORLINE::"):
            para(doc, line[len("AUTHORLINE::"):], align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(11))
        elif line.startswith("ABSTRACT::"):
            p = para(doc, "", size=Pt(9))
            r = p.add_run("Abstract—")
            r.bold, r.italic = True, True
            r.font.size = Pt(9)
            add_runs(p, line[len("ABSTRACT::"):], size=Pt(9), bold=True, italic=True)
        elif line.startswith("KEYTERMS::"):
            p = para(doc, "", size=Pt(9))
            r = p.add_run("Key Terms—")
            r.bold, r.italic = True, True
            r.font.size = Pt(9)
            add_runs(p, line[len("KEYTERMS::"):], size=Pt(9), italic=True)
        elif line.startswith("TABLE::"):
            cap = para(doc, line[len("TABLE::"):], align=WD_ALIGN_PARAGRAPH.CENTER,
                       size=Pt(8), italic=False)
            cap.paragraph_format.keep_with_next = True
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            rows = [rows[0]] + rows[2:]  # выкинуть разделитель ---
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for ri, row in enumerate(rows):
                for ci, val in enumerate(row):
                    set_cell(t.cell(ri, ci), val, Pt(8), bold=(ri == 0),
                             align=WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER)
            table_borders(t)
            # не рвать таблицу между страницами
            for row in t.rows[:-1]:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True
            ncols = len(rows[0])
            if ncols == 6:      # Table I: horizon/median/mean/t/Wilcoxon/%>0
                fix_widths(t, [0.52, 0.60, 0.56, 0.42, 0.76, 0.44])
            elif ncols == 4:    # Table II: hypothesis/raw/BH/verdict
                fix_widths(t, [1.22, 0.62, 0.68, 0.78])
        elif line.startswith("FIG::"):
            _, path, caption = line.split("::", 2)
            if not NOFIGS:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(str((MD.parent / path).resolve()), width=Inches(3.3))
                para(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, size=Pt(8))
        elif line.startswith("REF::"):
            if NOFIGS:
                continue
            para(doc, line[len("REF::"):], size=Pt(8),
                 align=WD_ALIGN_PARAGRAPH.LEFT)
        elif line.startswith("# "):
            text = line[2:].strip()
            if first_h1:
                para(doc, text, style="Title")
                first_h1 = False
            else:
                in_refs = text == "References"
                if NOFIGS and in_refs:
                    break
                para(doc, text, style="Heading 1")
        elif line.startswith("## "):
            para(doc, line[3:].strip(), style="Heading 2")
        else:
            para(doc, line, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    OUT_DIR.mkdir(exist_ok=True)
    doc.save(str(OUT))
    print("saved:", OUT)


if __name__ == "__main__":
    main()
