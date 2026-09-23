"""Сборка BakhtievSaid_figures.pptx из официального CJSJ-Figures-Template.pptx.

Структура шаблона (разобрана): 16:9 (10 x 5.625 in), титульный слайд +
по слайду на фигуру/таблицу: картинка/таблица по центру, caption 14pt внизу
(«Figure 1.\\t...» / «Table 1.\\t...»). Примерные слайды шаблона выкидываются,
содержимое берётся из paper/cjsj_submission.md (FIG::/TABLE:: блоки), чтобы
подписи и числа не расходились с текстом статьи.

Запуск: /tmp/cjsj-build/bin/python publish/build_figures_pptx.py
"""
import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
MD = ROOT / "paper" / "cjsj_submission.md"
TPL = Path.home() / "РЕСЕРЧ" / "repro_cjsj_v1" / "template" / "CJSJ-Figures-Template.pptx"
OUT = ROOT / "publish" / "cjsj_package" / "BakhtievSaid_figures.pptx"

PAPER_TITLE = ("Post-Listing Underperformance in Cryptocurrency Markets: "
               "Evidence from Every Binance USDT Listing, 2021-2026")
SUBTITLE = "Said Bakhtiev, Gymnasium of Aznakayevo, Republic of Tatarstan, Russia; 2026-2027"

CAP_W, CAP_H = Inches(9.32), Inches(0.63)
CAP_LEFT, CAP_TOP = Inches(0.34), Inches(4.83)
IMG_BOX_W, IMG_BOX_H = 6.75, 4.62  # inches, из разбора шаблона


def delete_slide(prs, slide):
    for sldId in list(prs.slides._sldIdLst):
        if prs.part.related_part(sldId.rId) == slide.part:
            prs.part.drop_rel(sldId.rId)
            prs.slides._sldIdLst.remove(sldId)
            return


def add_caption(slide, text):
    box = slide.shapes.add_textbox(CAP_LEFT, CAP_TOP, CAP_W, CAP_H)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    label, _, rest = text.partition("\t")
    r1 = p.add_run(); r1.text = label + "\t"
    r1.font.size = Pt(14); r1.font.name = "Times New Roman"; r1.font.bold = True
    r2 = p.add_run(); r2.text = rest
    r2.font.size = Pt(14); r2.font.name = "Times New Roman"


def add_picture_fit(slide, path):
    with Image.open(path) as im:
        w, h = im.size
    aspect = w / h
    bw, bh = IMG_BOX_W, IMG_BOX_H
    if bw / bh > aspect:
        bh = bh
        bw = bh * aspect
    else:
        bh = bw / aspect
    left = Inches((10.0 - bw) / 2)
    top = Inches((4.83 - bh) / 2)
    slide.shapes.add_picture(str(path), left, top, Inches(bw), Inches(bh))


def add_native_table(slide, rows):
    n_r, n_c = len(rows), len(rows[0])
    width = Inches(8.6)
    height = Inches(0.35 * n_r)
    left = Inches((10.0 - 8.6) / 2)
    top = Inches(max(0.4, (4.83 - 0.35 * n_r) / 2))
    gfx = slide.shapes.add_table(n_r, n_c, left, top, width, height)
    tbl = gfx.table
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.size = Pt(14)
                    r.font.name = "Times New Roman"
                    r.font.bold = (ri == 0)


def main():
    text = MD.read_text(encoding="utf-8")
    figs = re.findall(r"^FIG::(\.\./charts/\S+?)::(.+)$", text, re.M)
    blocks = re.findall(r"^TABLE::(.+)$\n((?:\|.*\n?)+)", text, re.M)

    prs = Presentation(str(TPL))
    slides = list(prs.slides)
    # титул: правим текст на месте
    title_slide = slides[0]
    phs = [sh for sh in title_slide.shapes if sh.has_text_frame]
    phs[0].text_frame.text = PAPER_TITLE
    phs[1].text_frame.text = SUBTITLE
    # примерные слайды выкинуть
    for s in slides[1:]:
        delete_slide(prs, s)

    blank = prs.slide_layouts[1]  # TITLE_AND_BODY; плейсхолдеры чистим ниже

    def fresh_slide():
        s = prs.slides.add_slide(blank)
        for sh in list(s.shapes):
            sh._element.getparent().remove(sh._element)
        return s

    for i, (path, cap) in enumerate(figs, start=1):
        s = fresh_slide()
        add_picture_fit(s, (MD.parent / path).resolve())
        cap = re.sub(r"^Fig\. \d+\.\s*", "", cap)
        add_caption(s, f"Figure {i}.\t{cap}")

    for i, (cap, tbl_md) in enumerate(blocks, start=1):
        s = fresh_slide()
        rows = []
        for ln in tbl_md.strip().splitlines():
            cells = [c.strip() for c in ln.strip().strip("|").split("|")]
            if set("".join(cells)) <= set(":- "):
                continue
            rows.append(cells)
        add_native_table(s, rows)
        cap = re.sub(r"^TABLE [IVX]+\.\s*", "", cap)
        add_caption(s, f"Table {i}.\t{cap}")

    OUT.parent.mkdir(exist_ok=True)
    prs.save(str(OUT))
    print("saved:", OUT, "slides:", len(list(prs.slides)))


if __name__ == "__main__":
    main()
